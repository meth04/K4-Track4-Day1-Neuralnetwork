"""results_table.py — HOÀN THIỆN. Lưu kết quả ra JSON rồi điền vào experiments.xlsx từ mẫu.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

FORMULA_COLS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}


def _jsonable(obj):
    """Chuyển numpy/tuple sang kiểu JSON được."""
    import numpy as np
    if isinstance(obj, dict):
        return {k: _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        return float(obj)
    if isinstance(obj, (np.bool_,)):
        return bool(obj)
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    return obj


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg + history + summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    payload = dict(cfg=result["cfg"], history=result["history"], summary=result["summary"])
    path = Path(results_dir) / f"{exp_id}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(_jsonable(payload), f, indent=2, ensure_ascii=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    out = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            out.append(json.load(f))
    return out


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc/eval_macro_f1).

    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.
    """
    cfg, s = result["cfg"], result["summary"]
    hidden = cfg.get("hidden", ())
    hidden_str = "-".join(map(str, hidden)) if not isinstance(hidden, str) else hidden
    row = dict(
        exp_id=cfg.get("exp_id"),
        group=cfg.get("group"),
        description=cfg.get("description"),
        loss=cfg.get("loss"),
        optimizer=cfg.get("optimizer"),
        lr=cfg.get("lr"),
        weight_decay=cfg.get("weight_decay"),
        batch=cfg.get("batch"),
        epochs=cfg.get("epochs"),
        hidden=hidden_str,
        dropout=cfg.get("dropout"),
        clip_norm=("none" if cfg.get("clip_norm") is None else cfg.get("clip_norm")),
        precision=cfg.get("precision"),
        init=cfg.get("init"),
        seed=cfg.get("seed"),
        step0_loss=s.get("step0_loss"),
        best_val_loss=s.get("best_val_loss"),
        best_epoch=s.get("best_epoch"),
        final_train_loss=s.get("final_train_loss"),
        final_val_loss=s.get("final_val_loss"),
        val_acc=s.get("val_acc"),
        val_macro_f1=s.get("val_macro_f1"),
        time_per_epoch_s=s.get("time_per_epoch_s"),
        peak_mem_MB=(None if s.get("peak_mem_MB") != s.get("peak_mem_MB") else s.get("peak_mem_MB")),
        diverged=bool(s.get("diverged")),
        eval_acc=(eval_scores or {}).get("eval_acc"),
        eval_macro_f1=(eval_scores or {}).get("eval_macro_f1"),
        figure_file=f"figures/{cfg.get('exp_id')}.png",
        notes=notes or cfg.get("description", ""),
    )
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               baseline_seed_ids: list[str] | None = None,
               summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    - Giữ nguyên công thức ở các cột cuối (AD..AG) và sheet Seeds/Summary.
    - baseline_seed_ids: ghi vào cột A của sheet Seeds (các lần chạy baseline khác seed).
    - summary_notes: {group: nhận xét} ghi vào cột H của sheet Summary.
    """
    wb = openpyxl.load_workbook(template_path)  # KHÔNG data_only=True (giữ công thức)
    ws = wb["Experiments"]

    header = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v:
            header[str(v).strip()] = c

    max_row = ws.max_row  # 61 (1 tiêu đề + 60 dòng)
    # xoá dữ liệu cũ ở các cột KHÔNG phải công thức (tránh sót dữ liệu khi chạy lại)
    for r in range(2, max_row + 1):
        for name, c in header.items():
            if name not in FORMULA_COLS:
                ws.cell(r, c).value = None

    for i, row in enumerate(rows):
        r = 2 + i
        if r > max_row:
            raise ValueError(f"có {len(rows)} dòng > {max_row - 1} dòng của mẫu; hãy tăng số dòng trong mẫu")
        for name, val in row.items():
            if name in FORMULA_COLS:
                continue
            c = header.get(name)
            if c is None:
                continue
            ws.cell(r, c).value = val

    # ---- sheet Seeds: ghi exp_id baseline theo seed (cột A), công thức B/C/D tự tính
    if baseline_seed_ids:
        ws_seeds = wb["Seeds"]
        for i, exp_id in enumerate(baseline_seed_ids):
            ws_seeds.cell(2 + i, 1).value = exp_id

    # ---- sheet Summary: nhận xét ngắn (cột H)
    if summary_notes:
        ws_sum = wb["Summary"]
        group_col = {str(ws_sum.cell(r, 1).value): r for r in range(2, ws_sum.max_row + 1)}
        for grp, note in summary_notes.items():
            r = group_col.get(grp)
            if r:
                ws_sum.cell(r, 8).value = note

    wb.save(out_path)
    print(f"đã ghi {out_path} với {len(rows)} dòng thí nghiệm")
